"""Reads the MasterView sheet of MASTER.xlsx into a list of ParamRow."""
from __future__ import annotations

from pathlib import Path

import openpyxl

from dictionary.models import ParamRow

SHEET_NAME = "MasterView"


def load_dictionary(xlsx_path: Path) -> list[ParamRow]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"Sheet '{SHEET_NAME}' not found in {xlsx_path}. "
                          f"Available sheets: {', '.join(wb.sheetnames)}")
    ws = wb[SHEET_NAME]
    rows = []
    for r in ws.iter_rows(min_row=2, max_col=12, values_only=True):
        (row_number, scope, category, subcategory, parameter, input_type, unit,
         aff_sub, aff_param, data_validation, tooltip, user_comment) = r
        if not parameter:
            continue
        rows.append(ParamRow(
            row_number=int(row_number),
            scope=(scope or "").strip(),
            category=(category or "").strip(),
            subcategory=(subcategory or "").strip(),
            parameter=str(parameter).strip(),
            input_type=(input_type or "").strip(),
            unit=(unit or "").strip() if unit else "",
            affected_subcategory=str(aff_sub).strip() if aff_sub else "",
            affected_parameter=str(aff_param).strip() if aff_param else "",
            data_validation=str(data_validation).strip() if data_validation else "",
            tooltip=str(tooltip).strip() if tooltip else "",
            user_comment=str(user_comment).strip() if user_comment else "",
        ))
    wb.close()
    return rows
